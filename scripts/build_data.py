"""Generate src/data/exercises.json from data/training_dictionary_v2.xlsx and
validate src/data/templates.json against it.

Usage:  python scripts/build_data.py
Re-run whenever you edit the spreadsheet's 动作库 Exercises sheet.
"""
import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "data" / "training_dictionary_v2.xlsx"
OUT = ROOT / "src" / "data" / "exercises.json"
TEMPLATES = ROOT / "src" / "data" / "templates.json"

# Category code -> rotation group shown in the app
GROUP = {
    "WU": "热身", "SQ": "腿", "HG": "腿", "SL": "腿", "HPU": "胸", "VPL": "背", "HPL": "背",
    "VPU": "肩", "SH": "肩", "BI": "二头", "TR": "三头", "CO": "核心", "PW": "全身", "BL": "平衡",
}
LOWER = {"SQ", "HG", "SL"}

# How an exercise's load is recorded
LEVEL = {"VPL01"}  # assisted pull-up machine: weight field = assist level (lower is harder)
TIMED = {"WU01", "WU02", "SQ07", "SQ12", "CO08", "CO12", "BL01", "BL04"}
BODYWEIGHT_EQUIP = ("自重", "TRX", "单杠", "引体架", "罗马椅", "健腹轮", "弹力带", "BOSU")


def load_type(eid: str, equip: str) -> str:
    if eid in LEVEL:
        return "level"
    if eid in TIMED:
        return "sec"
    if any(equip.startswith(b) for b in BODYWEIGHT_EQUIP) and "/" not in equip:
        return "bw"
    return "lb"


def increment(cat: str, equip: str) -> int:
    heavy = any(k in equip for k in ("杠铃", "器械", "雪橇"))
    return 10 if cat in LOWER and heavy else 5


def main() -> None:
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["动作库 Exercises"]
    header = [c.value for c in ws[1]]
    col = {name: i for i, name in enumerate(header)}

    def get(row, name):
        v = row[col[name]]
        return "" if v is None else str(v).strip()

    exercises = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        eid = get(row, "ID")
        cat = get(row, "分类代码")
        equip = get(row, "器材 Equipment")
        legacy = [s for s in get(row, "历史记录（来源#组合: 重量×次数）").split("\n") if s]
        exercises.append({
            "id": eid,
            "cn": get(row, "中文名"),
            "en": get(row, "English name"),
            "cat": cat,
            "catName": get(row, "分类"),
            "size": get(row, "肌群大小"),
            "group": GROUP[cat],
            "equip": equip,
            "home": get(row, "可在家 Home?") == "Y",
            "homeNote": get(row, "居家用法/替代"),
            "muscles": get(row, "主要肌群"),
            "cue": get(row, "动作要点"),
            "legacy": legacy,
            "load": load_type(eid, equip),
            "inc": increment(cat, equip),
        })

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(exercises, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"wrote {len(exercises)} exercises -> {OUT.relative_to(ROOT)}")

    ids = {e["id"] for e in exercises}
    home_ids = {e["id"] for e in exercises if e["home"]}
    errors = []
    for t in json.loads(TEMPLATES.read_text(encoding="utf-8")):
        for it in t["items"]:
            if it["ex"] not in ids:
                errors.append(f"{t['key']}: unknown exercise {it['ex']}")
            elif t["loc"] == "home" and it["ex"] not in home_ids:
                errors.append(f"{t['key']}: {it['ex']} is not marked home-able")
    if errors:
        print("\n".join(errors))
        sys.exit(1)
    print("templates.json OK")


if __name__ == "__main__":
    main()
